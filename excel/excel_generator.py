from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from io import BytesIO

class Excel:
    @staticmethod
    def generar_excel_estilo_laboratorio(muestras):
        wb = Workbook()
        ws = wb.active
        ws.title = "Resultados"

        # Estilos
        titulo = Font(bold=True, size=14)
        encabezado = Font(bold=True)
        centro = Alignment(horizontal="center", vertical="center")
        borde = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        # Título
        ws.merge_cells("A1:O1")
        ws["A1"] = "RESULTADOS DE MUESTRAS DE AGUA"
        ws["A1"].font = titulo
        ws["A1"].alignment = centro

        # Encabezados
        headers = [
            "Código", "Fecha", "Hora", "Quién Muestra", "Dispositivo",
            "Fuente", "Código Red", "Tipo Agua", "pH", "Cloro",
            "Temperatura", "Observaciones",
            "Fecha Recepción", "Hora Recepción", "Temperatura Recepción"
        ]
        
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=col, value=h)
            c.font = encabezado
            c.alignment = centro
            c.border = borde
            ws.column_dimensions[c.column_letter].width = 18

        # Datos
        fila = 4
        for m in muestras:
            fila_data = [
                m.get("Código"),
                m.get("Fecha"),
                m.get("Hora"),
                m.get("Quién Muestra"),
                m.get("Dispositivo"),
                m.get("Fuente"),
                m.get("Código Red"),
                m.get("Tipo Agua"),
                m.get("pH"),
                m.get("Cloro"),
                m.get("Temperatura"),
                m.get("Observaciones"),
                m.get("Fecha Recepción"),
                m.get("Hora Recepción"),
                m.get("Temperatura Recepción")
            ]

            for col, val in enumerate(fila_data, 1):
                c = ws.cell(row=fila, column=col, value=val)
                c.border = borde
                c.alignment = centro

            fila += 1

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()